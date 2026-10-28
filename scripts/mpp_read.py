"""Baca XML MS Project (MSPDI, dari mpp_export.ps1) -> senarai aktiviti dalam Excel + JSON.
Guna:  py -3.12 scripts/mpp_read.py "output/programme/<nama>.xml" [--aras 2] [--cari teks]
Hasil: output/programme/<nama>.xlsx  dan  output/programme/<nama>.json
       Cetak aktiviti hingga --aras (default 2), atau yang padan --cari.
"""
import argparse
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

NS = {"p": "http://schemas.microsoft.com/project"}
LINK = {"0": "FF", "1": "FS", "2": "SF", "3": "SS"}

ap = argparse.ArgumentParser()
ap.add_argument("xml")
ap.add_argument("--aras", type=int, default=2, help="aras WBS maksimum untuk dicetak")
ap.add_argument("--cari", help="cetak aktiviti yang namanya mengandungi teks ini")
args = ap.parse_args()

src = Path(args.xml)
root = ET.parse(src).getroot()
mins_per_day = int(root.findtext("p:MinutesPerDay", "480", NS))


def days(iso):
    """PT123H30M0S -> hari kerja."""
    m = re.fullmatch(r"PT(\d+)H(\d+)M(\d+)S", iso or "")
    if not m:
        return None
    h, mi, _ = map(int, m.groups())
    return round((h * 60 + mi) / mins_per_day, 1)


def date(s):
    return s[:10] if s else ""


tasks, by_uid = [], {}
for t in root.findall("p:Tasks/p:Task", NS):
    get = lambda tag: t.findtext(f"p:{tag}", "", NS)
    if not get("Name"):
        continue
    preds = [
        get_uid + LINK.get(pl.findtext("p:Type", "1", NS), "FS")
        for pl in t.findall("p:PredecessorLink", NS)
        for get_uid in [pl.findtext("p:PredecessorUID", "", NS)]
    ]
    row = {
        "uid": get("UID"),
        "id": get("ID"),
        "wbs": get("WBS") or get("OutlineNumber"),
        "aras": int(get("OutlineLevel") or 0),
        "nama": get("Name").strip(),
        "tempoh_hari": days(get("Duration")),
        "mula": date(get("Start")),
        "siap": date(get("Finish")),
        "peratus_siap": int(get("PercentComplete") or 0),
        "ringkasan": get("Summary") == "1",
        "milestone": get("Milestone") == "1",
        "kritikal": get("Critical") == "1",
        "pendahulu_uid": preds,
    }
    tasks.append(row)
    by_uid[row["uid"]] = row

# Tukar UID pendahulu kepada ID yang dilihat dalam MS Project
for r in tasks:
    r["pendahulu"] = ", ".join(
        by_uid[p[:-2]]["id"] + (p[-2:] if p[-2:] != "FS" else "")
        for p in r.pop("pendahulu_uid") if p[:-2] in by_uid
    )

out_json = src.with_suffix(".json")
out_json.write_text(json.dumps(tasks, ensure_ascii=False, indent=1), encoding="utf-8")

wb = Workbook()
ws = wb.active
ws.title = "Aktiviti"
head = ["ID", "WBS", "Nama aktiviti", "Tempoh (hari)", "Mula", "Siap", "% Siap", "Pendahulu", "Milestone", "Kritikal"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for r in tasks:
    ws.append([int(r["id"] or 0), r["wbs"], r["nama"], r["tempoh_hari"], r["mula"], r["siap"],
               r["peratus_siap"], r["pendahulu"], "Ya" if r["milestone"] else "", "Ya" if r["kritikal"] else ""])
    cell = ws.cell(ws.max_row, 3)
    cell.alignment = Alignment(indent=max(r["aras"] - 1, 0))
    if r["ringkasan"]:
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
for col, w in zip("ABCDEFGHIJ", [6, 12, 70, 12, 12, 12, 8, 18, 10, 9]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
out_xlsx = src.with_suffix(".xlsx")
wb.save(out_xlsx)

print(f"{len(tasks)} aktiviti, {sum(r['milestone'] for r in tasks)} milestone, "
      f"{sum(r['kritikal'] for r in tasks)} kritikal")
print(f"Excel: {out_xlsx}\nJSON:  {out_json}\n")
pick = [r for r in tasks if args.cari.lower() in r["nama"].lower()] if args.cari \
    else [r for r in tasks if r["aras"] <= args.aras]
for r in pick:
    print(f"{r['wbs']:12s} {'  ' * max(r['aras'] - 1, 0)}{r['nama'][:70]:70s} "
          f"{r['tempoh_hari'] or 0:6.0f}h  {r['mula']} - {r['siap']}")
