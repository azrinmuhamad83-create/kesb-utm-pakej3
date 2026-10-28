"""Jana Senarai Induk Kelulusan Bahan & Submittal (Excel) untuk dikemas kini oleh Kerani Tapak.

Guna:  py -3.12 scripts/build_daftar_kelulusan.py          (cipta fail jika belum ada)
       py -3.12 scripts/build_daftar_kelulusan.py --paksa  (tulis ganti – SEMUA kemas kini Kerani hilang)
Hasil: output/Daftar-Kelulusan-Bahan.xlsx

Fail ini rekod hidup (CP-QC Seksyen 6.1; KESB-QF-21). Selepas dicipta, ia dikemas kini terus dalam Excel;
skrip ini tidak menulis ganti fail sedia ada tanpa --paksa.
Senarai item: CP-QC Seksyen 6.2 dan aktiviti kelulusan bahan dalam ITP. Tarikh perlu: Master Programme REV00
(mula kerja Fasa 1); shop drawing 1 bulan sebelum kerja (Prelim 11.0); method statement 14 hari selepas SST (Prelim 3.0).
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("output/Daftar-Kelulusan-Bahan.xlsx")
HEAD = PatternFill("solid", fgColor="1F4E78")
TRADE = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS = ["Belum dikemukakan", "Dikemukakan", "Lulus", "Lulus bersyarat", "Ditolak – kemuka semula", "Tidak berkaitan"]

SB, SB_SD = (date(2026, 9, 1), "Seni bina Fasa 1"), (date(2026, 8, 1), "Shop drawing – seni bina Fasa 1")
MK, MK_SD = (date(2026, 9, 3), "Mekanikal Fasa 1"), (date(2026, 8, 3), "Shop drawing – mekanikal Fasa 1")
EL, EL_SD = (date(2027, 4, 5), "Elektrik FKOM"), (date(2027, 3, 5), "Shop drawing – elektrik FKOM")
PR = (date(2027, 5, 15), "Perabot (BQ C)")
P3 = (date(2026, 8, 24), "Prelim 3.0 – 14 hari selepas SST")

# (trade, item, jenis, rujukan, (tarikh perlu, sebab))
ITEMS = [
    ("Dokumen Prelim 3.0", "Programme of Works CPM (dengan CV perancang)", "Dokumen", "Prelim 3.0 (i)", P3),
    ("Dokumen Prelim 3.0", "Carta organisasi & CV personel utama", "Dokumen", "Prelim 3.0 (iii)", P3),
    ("Dokumen Prelim 3.0", "Jadual penyelenggaraan & kalibrasi serta pemeriksaan & ujian loji (Lampiran E)", "Dokumen", "Prelim 3.0 (iv)", P3),
    ("Dokumen Prelim 3.0", "Senarai subkontraktor & Program Kerja", "Dokumen", "SST para 12", P3),
    ("Dinding & partition", "Cement board partition 2 x 12 mm + kerangka M.S hollow + fibrewool", "Bahan", "S-JK04; BQ B", SB),
    ("Dinding & partition", "Partition papan gipsum 2 x 12 mm + kerangka aluminium + fibrewool", "Bahan", "S-JK04; BQ B", SB),
    ("Dinding & partition", "Papan tulis kaca tempered 9 mm + marker tray aluminium", "Bahan", "S-JK04", SB),
    ("Dinding & partition", "Glass wall panel kaca 5 mm + kerangka aluminium + frosted sticker", "Bahan", "S-JK04", SB),
    ("Dinding & partition", "Panel akustik dinding (MDF berlubang / fabrik)", "Sampel", "S-JK04", SB),
    ("Dinding & partition", "Bata, plaster & exmet (dinding bata baharu)", "Bahan", "S-JK04; BQ B", SB),
    ("Dinding & partition", "Laminated polyester blockboard 18 mm (backdrop / panel)", "Sampel", "S-JK04", SB),
    ("Dinding & partition", "Skirting aluminium", "Bahan", "S-JK04", SB),
    ("Cat", "Sistem cat dalaman (wall sealer + emulsion)", "Bahan", "S-JK04; BQ B", SB),
    ("Cat", "Sistem cat luaran (wall sealer + weatherproof)", "Bahan", "S-JK04; BQ B", SB),
    ("Cat", "Cat anti-karat & gloss (gril, kerangka, tingkap)", "Bahan", "S-JK03–S-JK04", SB),
    ("Cat", "Warna & mock-up cat (feature wall)", "Sampel", "S-JK04; Prelim 12.0–13.0", SB),
    ("Lantai", "Jubin porselin 300 x 600 mm (F1) + adhesif + grout", "Sampel", "S-JK03 (F1); BQ B", SB),
    ("Lantai", "Plat expansion joint aluminium & skirting jubin 100 mm", "Bahan", "S-JK03; S-JK05", SB),
    ("Lantai", "Sistem epoxy F2 (epoxy mortar 4 mm, primer, self-levelling 1 mm, PU top coat UV)", "Bahan", "S-JK03 (F2); BQ B", SB),
    ("Lantai", "Bahan screed & ejen ikatan", "Bahan", "S-JK03", SB),
    ("Lantai", "Bahan kimia pembersih / penggilap jubin (F6)", "Bahan", "S-JK03 (F6); BQ B", SB),
    ("Lantai", "Karpet jubin (F8) / jubin terracotta (F5) – jika digunakan", "Sampel", "S-JK03–S-JK04", SB),
    ("Kalis air", "Sistem polyurea F9 (Purea Prime + Quartz 0.5, Purea-SA, Spartic-HA)", "Bahan", "S-JK04 (F9)", SB),
    ("Kalis air", "Sijil pemasang polyurea bertauliah pengilang", "Sijil", "S-JK04; Prelim 13.0", SB),
    ("Siling", "Papan gipsum 600 x 600 x 9 mm & tee aluminium", "Bahan", "S-JK05 (C1–C5a)", SB),
    ("Siling", "Panel akustik siling & plasterglass", "Sampel", "S-JK05", SB),
    ("Pintu & gril", "Pintu kaca + kerangka aluminium (Brown Anodized)", "Bahan", "S-JK03", SB),
    ("Pintu & gril", "Pintu kayu solid + kerangka logam", "Bahan", "S-JK03", SB),
    ("Pintu & gril", "Pintu kalis api 1 jam (D14 / D14a)", "Shop drawing", "S-JK03", SB_SD),
    ("Pintu & gril", "Sijil BOMBA pintu kalis api", "Sijil", "S-JK03", SB),
    ("Pintu & gril", "Pintu sliding kaca & pintu louvre aluminium", "Bahan", "S-JK03", SB),
    ("Pintu & gril", "Roller shutter (manual / bermotor)", "Shop drawing", "S-JK03", SB_SD),
    ("Pintu & gril", "Gril M.S (G1 / G1a)", "Bahan", "S-JK03", SB),
    ("Pintu & gril", "Ironmongery (lockset, door closer, engsel, pemegang, flush bolt)", "Sampel", "S-JK03", SB),
    ("Tingkap", "Tingkap aluminium Brown Anodized + kaca 5 mm bertinted (W6–W7a)", "Shop drawing", "S-JK05", SB_SD),
    ("Tingkap", "Tinted sticker tingkap", "Sampel", "S-JK04", SB),
    ("Bumbung, logam & konkrit", "Bumbung ACP 4 mm + struktur keluli + flashing (RF1, Blok P23)", "Shop drawing", "S-JK05; BQ B", SB_SD),
    ("Bumbung, logam & konkrit", "Railing RL1 M.S", "Bahan", "S-JK05", SB),
    ("Bumbung, logam & konkrit", "Konkrit Gred 25 & BRC A8 (papak F7, longkang)", "Bahan", "S-JK03; S-JK05", SB),
    ("Lab modular & perabot bina dalam", "Shop drawing lab modular, perabot terbina & pantri", "Shop drawing", "S-JK06–S-JK18", SB_SD),
    ("Lab modular & perabot bina dalam", "Worktop epoxy resin / ESD laminate / solid surface", "Sampel", "S-JK06–S-JK18", SB),
    ("Lab modular & perabot bina dalam", "Kabinet & hardware", "Sampel", "S-JK06–S-JK18", SB),
    ("Petanda", "Sampel petanda SP01 / SP02 & senarai teks", "Sampel", "S-JK19", SB),
    ("Perabot lepas", "Kerusi", "Sampel", "P-JP01–P-JP14; BQ C", PR),
    ("Perabot lepas", "Meja", "Sampel", "P-JP; BQ C", PR),
    ("Perabot lepas", "Almari & kabinet", "Sampel", "P-JP11–P-JP13; BQ C", PR),
    ("Perabot lepas", "Roller blind", "Sampel", "P-JP; BQ C", PR),
    ("Perabot lepas", "Perabot upholsteri", "Sampel", "P-JP; BQ C", PR),
    ("Hawa dingin & pengudaraan", "Unit hawa dingin split R32 (ceiling cassette, wall mounted, ceiling exposed)", "Bahan", "BQ D; M-001", MK),
    ("Hawa dingin & pengudaraan", "Paip tembaga, penebat Armaflex & paip kondensat PVC", "Bahan", "BQ D; M-001", MK),
    ("Hawa dingin & pengudaraan", "Starter panel metalclad", "Bahan", "BQ D", MK),
    ("Hawa dingin & pengudaraan", "AHU double skin 150,000 btu/hr (Blok E04)", "Bahan", "BQ D", MK),
    ("Hawa dingin & pengudaraan", "Laluan paip & lokasi unit hawa dingin", "Shop drawing", "BQ D; M-FK / M-FKM", MK_SD),
    ("Hawa dingin & pengudaraan", "Exhaust fan, axial fan 200 mm & ducting", "Bahan", "BQ D", MK),
    ("Pencegah kebakaran", "Detector (asap / haba)", "Bahan", "BQ D", MK),
    ("Pencegah kebakaran", "Lampu kecemasan & tanda KELUAR", "Bahan", "BQ D", MK),
    ("Elektrik", "DB, MCCB / MCB / RCCB & SPD", "Bahan", "BQ E", EL),
    ("Elektrik", "Skematik DB (termasuk ruang tanpa skematik)", "Shop drawing", "BQ E; E-FK / E-FKM", EL_SD),
    ("Elektrik", "Kabel sub-utama, pendawaian & konduit / trunking GS", "Bahan", "BQ E", EL),
    ("Elektrik", "Lampu LED", "Sampel", "BQ E", EL),
    ("Elektrik", "Suis, soket 13A & kipas", "Bahan", "BQ E", EL),
    ("Elektrik", "Sijil Orang Kompeten PW4 + Endorsan Pengujian", "Sijil", "BQ E nota IV", EL),
    ("ICT", "Kabel CAT6 / CAT6A, panel tamatan & rak", "Bahan", "BQ E", EL),
    ("ICT", "WiFi 7 AP & network switch", "Bahan", "BQ E", EL),
    ("Audio visual", "Projektor 4500 lumen, skrin bermotor & TV 75\"", "Bahan", "BQ E", EL),
    ("Audio visual", "Speaker, mixer, amplifier, matrix & mikrofon UHF", "Bahan", "BQ E", EL),
    ("Audio visual", "Susun atur AV", "Shop drawing", "BQ E", EL_SD),
]


def method_statements():
    rows = []
    for f in sorted(Path("output/method-statement").glob("MS-*.md")):
        rows.append(("Method statement", f.stem, "Method statement", "Prelim 3.0 (ii)", P3))
    return rows


def main():
    if OUT.exists() and "--paksa" not in sys.argv:
        print(f"{OUT} sudah ada – tidak ditulis ganti (kemas kini Kerani dikekalkan). Guna --paksa untuk cipta semula.")
        return
    items = ITEMS[:4] + method_statements() + ITEMS[4:]
    wb = Workbook()

    ws = wb.active
    ws.title = "Senarai Induk"
    ws.append(["Senarai Induk Kelulusan Bahan & Submittal – UTM Baikpulih Makmal P&P, Pakej 3 (KESB/UTM-P3)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Dikemas kini oleh Kerani Tapak. Status rasmi hanya selepas borang KESB-QF-01 / surat ditandatangani S.O. "
               "Baris merah = tarikh perlu telah lepas dan belum lulus."])
    head = ["No", "Trade", "Item", "Jenis", "Rujukan (BQ / lukisan)", "Perlu lulus sebelum", "Asas tarikh perlu",
            "No. MAR / rujukan", "Tarikh dikemukakan", "Status", "Tarikh keputusan", "Diputuskan oleh",
            "Bukti (QF-01 bertandatangan / surat / e-mel / WhatsApp)", "Catatan"]
    ws.append(head)
    widths = [5, 22, 48, 13, 22, 12, 24, 16, 12, 20, 12, 18, 30, 30]
    for cell, w in zip(ws[3], widths):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    for i, (trade, item, jenis, ruj, (tarikh, asas)) in enumerate(items, start=1):
        ws.append([i, trade, item, jenis, ruj, tarikh, asas, "", None, "Belum dikemukakan", None, "", "", ""])
        for c in ws[ws.max_row]:
            c.alignment = WRAP
            c.border = BORDER
        for col in ("F", "I", "K"):
            ws[f"{col}{ws.max_row}"].number_format = "DD/MM/YYYY"
    # Baris kosong untuk item tambahan
    for i in range(len(items) + 1, len(items) + 21):
        ws.append([i] + [""] * 13)
        for c in ws[ws.max_row]:
            c.alignment = WRAP
            c.border = BORDER
        for col in ("F", "I", "K"):
            ws[f"{col}{ws.max_row}"].number_format = "DD/MM/YYYY"
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J4:J{last}")
    rng = f"A4:N{last}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND($F4<>"",$F4<TODAY(),$J4<>"Lulus",$J4<>"Lulus bersyarat",$J4<>"Tidak berkaitan")'],
        fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR($J4="Lulus",$J4="Lulus bersyarat")'],
                                                   fill=PatternFill("solid", fgColor="E2EFDA")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$J4="Dikemukakan"'], fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.freeze_panes = "D4"
    ws.auto_filter.ref = f"A3:N{last}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "3:3"

    # Ringkasan ikut trade & status (formula – dikemas kini sendiri)
    rs = wb.create_sheet("Ringkasan")
    rs.append(["Ringkasan status kelulusan (dikira daripada sheet Senarai Induk)"])
    rs["A1"].font = Font(bold=True, size=13)
    rs.append([])
    rs.append(["Trade"] + STATUS + ["Jumlah", "Lewat (belum lulus, tarikh lepas)"])
    for c in rs[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, horizontal="center")
        c.border = BORDER
    rs.column_dimensions["A"].width = 32
    for col in "BCDEFGHI":
        rs.column_dimensions[col].width = 14
    trades = list(dict.fromkeys(t[0] for t in items))
    S = f"'Senarai Induk'!$B$4:$B${last}"
    J = f"'Senarai Induk'!$J$4:$J${last}"
    F = f"'Senarai Induk'!$F$4:$F${last}"
    for t in trades + ["JUMLAH"]:
        r = rs.max_row + 1
        if t == "JUMLAH":
            row = ["JUMLAH"] + [f"=SUM({chr(66 + k)}4:{chr(66 + k)}{r - 1})" for k in range(len(STATUS) + 2)]
        else:
            row = [t] + [f'=COUNTIFS({S},$A{r},{J},"{s}")' for s in STATUS]
            row += [f"=COUNTIF({S},$A{r})",
                    f'=SUMPRODUCT(({S}=$A{r})*({F}<>"")*({F}<TODAY())*({J}<>"Lulus")*({J}<>"Lulus bersyarat")*({J}<>"Tidak berkaitan"))']
        rs.append(row)
        for c in rs[rs.max_row]:
            c.border = BORDER
            if t == "JUMLAH":
                c.font = Font(bold=True)
                c.fill = TRADE

    ar = wb.create_sheet("Arahan")
    for line in [
        "Arahan kemas kini (Kerani Tapak)",
        "",
        "1. Setiap kali MAR / shop drawing / method statement dihantar kepada S.O.: isi No. MAR, Tarikh dikemukakan, dan tukar Status kepada 'Dikemukakan'.",
        "2. Bila keputusan S.O. diterima: isi Tarikh keputusan, Diputuskan oleh, dan tukar Status kepada 'Lulus', 'Lulus bersyarat' atau 'Ditolak – kemuka semula'.",
        "3. Lajur Bukti: catat dokumen sokongan – borang KESB-QF-01 bertandatangan S.O., no. surat, e-mel, atau 'WhatsApp <nama>, <tarikh>' (tangkap layar disimpan).",
        "4. Kelulusan melalui WhatsApp sahaja direkod 'Lulus bersyarat' dengan catatan 'menunggu QF-01 bertandatangan' sehingga dokumen bertulis diterima.",
        "5. Item yang tiada dalam senarai: tambah pada baris kosong di bawah. Jangan padam baris; tandakan 'Tidak berkaitan' jika item tidak digunakan.",
        "6. Baris merah = tarikh perlu telah lepas dan belum lulus – maklumkan Pengurus Projek.",
        "7. Kuning = sedang menunggu keputusan S.O.; hijau = lulus.",
        "8. Tiada pesanan atau fabrikasi sebelum Status 'Lulus' / 'Lulus bersyarat' (CP-QC Seksyen 6.1).",
    ]:
        ar.append([line])
    ar["A1"].font = Font(bold=True, size=13)
    ar.column_dimensions["A"].width = 140

    wb.save(OUT)
    print(f"Siap: {OUT}  ({len(items)} item, {len(trades)} kumpulan)")


if __name__ == "__main__":
    main()
