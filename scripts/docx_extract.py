"""BQ / dokumen Word (.docx) -> teks biasa (perenggan + jadual).
Guna:  py -3.12 scripts/docx_extract.py docs/kontrak
Hasil: output/teks/<nama>.txt
"""
import sys
from pathlib import Path
import docx

src_dir = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/kontrak")
out_dir = Path("output/teks")
out_dir.mkdir(parents=True, exist_ok=True)

for f in sorted(src_dir.glob("*.docx")):
    d = docx.Document(f)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            cells = []
            for c in r.cells:
                x = c.text.strip().replace("\n", " ")
                if not cells or cells[-1] != x:
                    cells.append(x)
            if any(cells):
                lines.append(" | ".join(cells))
    (out_dir / f"{f.stem}.txt").write_text("\n".join(lines), encoding="utf-8")
    print(f"{f.name}: {len(lines)} baris, {len(d.tables)} jadual")
