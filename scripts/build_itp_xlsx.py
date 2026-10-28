"""Jana ITP.xlsx (ITP + daftar ujian bahan + daftar NCR + daftar dokumen) dari output/itp.json.

Guna:  py -3.12 scripts/build_itp_xlsx.py
Input: output/itp.json  (ditulis oleh skill quality-plan)
Hasil: output/ITP.xlsx

Format itp.json:
{
  "projek": "Nama projek",
  "itp": [
    {"trade": "Konkrit", "aktiviti": [
      {"aktiviti": "Ujian slump", "kriteria": "75 +/- 25 mm", "rujukan": "Spec B.3.4",
       "kekerapan": "Setiap lori", "rekod": "Borang QC-05", "titik": "W"}
    ]}
  ],
  "ujian_bahan": [
    {"bahan": "Konkrit", "ujian": "Kiub 7 & 28 hari", "standard": "MS EN 12390-3",
     "kekerapan": "1 set / 50 m3", "makmal": "[PERLU SAHKAN]"}
  ]
}
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
TRADE_FILL = PatternFill("solid", fgColor="DDEBF7")
HOLD_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, project, headers, widths):
    ws = wb.create_sheet(title)
    ws.append([f"{project} - {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(headers)
    for col, (cell, w) in enumerate(zip(ws[3], widths), start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    return ws


def style_row(ws, fill=None):
    for cell in ws[ws.max_row]:
        cell.alignment = WRAP
        cell.border = BORDER
        if fill:
            cell.fill = fill


def main():
    data = json.loads(Path("output/itp.json").read_text(encoding="utf-8"))
    project = data.get("projek", "Projek")
    wb = Workbook()
    wb.remove(wb.active)

    ws = sheet(wb, "ITP", project,
               ["No", "Aktiviti / Pemeriksaan", "Kriteria Penerimaan", "Rujukan Spec / Drawing",
                "Kekerapan", "Rekod / Borang", "H/W/R", "Kontraktor", "SO / Perunding", "Catatan"],
               [6, 40, 35, 22, 18, 20, 8, 14, 14, 20])
    for t_no, trade in enumerate(data.get("itp", []), start=1):
        ws.append([t_no, trade["trade"].upper()])
        style_row(ws, TRADE_FILL)
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for a_no, a in enumerate(trade.get("aktiviti", []), start=1):
            point = a.get("titik", "R").upper()
            ws.append([f"{t_no}.{a_no}", a.get("aktiviti", ""), a.get("kriteria", ""), a.get("rujukan", ""),
                       a.get("kekerapan", ""), a.get("rekod", ""), point, "", "", ""])
            style_row(ws, HOLD_FILL if point == "H" else None)
            ws.cell(ws.max_row, 7).alignment = Alignment(horizontal="center", vertical="top")
    ws.append([])
    ws.append(["", "H = Hold point (kerja tidak boleh diteruskan tanpa kelulusan bertulis S.O. atau pihak berkuasa yang dinyatakan dalam kriteria)"])
    ws.append(["", "W = Witness point (SO dimaklumkan dan boleh hadir)"])
    ws.append(["", "R = Review (semakan rekod / dokumen, atau permit dalaman oleh pihak yang dinyatakan dalam kriteria)"])

    ws = sheet(wb, "Ujian Bahan", project,
               ["No", "Bahan", "Ujian", "Standard", "Kekerapan", "Makmal", "Status"],
               [6, 25, 35, 22, 22, 20, 14])
    for i, u in enumerate(data.get("ujian_bahan", []), start=1):
        ws.append([i, u.get("bahan", ""), u.get("ujian", ""), u.get("standard", ""),
                   u.get("kekerapan", ""), u.get("makmal", ""), ""])
        style_row(ws)

    ws = sheet(wb, "Daftar NCR", project,
               ["No NCR", "Tarikh", "Lokasi / Elemen", "Butiran Ketidakakuran", "Punca", "Tindakan Pembetulan",
                "Tanggungjawab", "Tarikh Sasaran", "Tarikh Tutup", "Status"],
               [12, 12, 20, 35, 25, 35, 16, 14, 14, 10])
    for _ in range(20):
        ws.append([""] * 10)
        style_row(ws)

    ws = sheet(wb, "Daftar Dokumen", project,
               ["No Dokumen", "Tajuk", "Jenis", "Revisi", "Tarikh", "Status", "Diedar Kepada"],
               [18, 40, 16, 8, 12, 14, 25])
    for _ in range(20):
        ws.append([""] * 7)
        style_row(ws)

    # Koordinasi lokasi (CP-QC Seksyen 8.4): data ruang dari output/lokasi.json jika ada
    lok_path = Path("output/lokasi.json")
    lok = json.loads(lok_path.read_text(encoding="utf-8")) if lok_path.exists() else {"lokasi": []}
    lalai = lok.get("lalai", {})
    ws = sheet(wb, "Daftar Lokasi", project,
               ["No", "Fakulti", "Blok", "Ruang (ikut BQ B)", "Aras", "Item BQ", "Fasa", "Tarikh mula",
                "Tarikh siap", "Wakil pengguna (nama / jawatan)", "No. telefon wakil pengguna",
                "Dalam group WhatsApp (Y/T)", "Penyelia KESB", "Subkon seni bina", "Subkon elektrik",
                "Status", "Catatan"],
               [5, 8, 7, 45, 6, 9, 6, 11, 11, 26, 16, 12, 22, 16, 14, 12, 20])
    for i, r in enumerate(lok.get("lokasi", []), start=1):
        ws.append([i, r.get("fakulti", ""), r.get("blok", ""), r.get("ruang", ""), r.get("aras", ""),
                   r.get("item", ""), "", "", "", "", "", "", lalai.get("penyelia_kesb", ""),
                   lalai.get("subkon_senibina", ""), lalai.get("subkon_elektrik", ""), "", ""])
        style_row(ws)
    ws.freeze_panes = "E4"

    ws = sheet(wb, "Log Isu Pengguna", project,
               ["No. rujukan", "Tarikh / masa", "Jenis (aduan / henti kerja)", "Saluran (WhatsApp / telefon / surat / lisan)",
                "Blok / ruang", "Dilapor oleh (nama / jawatan)", "Isu / aduan / sebab", "Kerja & bil. pekerja terjejas",
                "Masa henti – sambung", "Pengesahan S.O. (rasmi / tidak rasmi, no. surat)", "Tindakan", "Tanggungjawab",
                "No. NCR", "Tarikh sasaran", "Tarikh selesai", "Status"],
               [11, 14, 12, 16, 20, 22, 32, 22, 16, 22, 30, 16, 12, 13, 13, 10])
    # Rekod isu kekal dalam output/log_isu.json (diisi oleh bot Telegram atau secara manual);
    # baris kosong bernombor ditambah selepas rekod terakhir untuk diisi tangan.
    LOG_KEYS = ["no", "tarikh_masa", "jenis", "saluran", "lokasi", "dilapor_oleh", "isu", "terjejas",
                "henti_sambung", "pengesahan_so", "tindakan", "tanggungjawab", "no_ncr", "tarikh_sasaran",
                "tarikh_selesai", "status"]
    log_path = Path("output/log_isu.json")
    log = json.loads(log_path.read_text(encoding="utf-8")).get("isu", []) if log_path.exists() else []
    for r in log:
        ws.append([r.get(k, "") for k in LOG_KEYS])
        style_row(ws)
    for i in range(len(log) + 1, len(log) + 26):
        ws.append([f"ISU-{i:03d}"] + [""] * 15)
        style_row(ws)

    out = Path("output/ITP.xlsx")
    wb.save(out)
    n = sum(len(t.get("aktiviti", [])) for t in data.get("itp", []))
    print(f"Siap: {out}  ({len(data.get('itp', []))} trade, {n} aktiviti)")


if __name__ == "__main__":
    main()
