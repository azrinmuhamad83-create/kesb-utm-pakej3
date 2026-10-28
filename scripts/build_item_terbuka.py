"""Kumpul semua item [PERLU SAHKAN] daripada dokumen projek ke dalam Daftar Item Terbuka (Excel).

Guna:  py -3.12 scripts/build_item_terbuka.py
Hasil: output/Daftar-Item-Terbuka.xlsx  (Lampiran J CP-QC)

Sumber: output/CP-QC-draf.md, output/itp.json, output/method-statement/MS-*.md, scripts/build_borang.py.
Item dengan teks sama digabung; semua lokasinya disenaraikan. Lampiran E diringkaskan kepada 3 item.
Jana semula selepas dokumen dikemas kini – item yang telah disahkan (tanda dibuang) akan hilang dari daftar.
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TAG = re.compile(r"\[PERLU SAHKAN:?\s*([^\]]*)\]")

# (kategori, kata kunci, dapatkan daripada, pemilik)
KATEGORI = [
    ("Kontrak & pentadbiran", r"SST|LAD|DLP|Lampiran A|PWD|klausa|tuntutan|retention|QLASSIC|CIS 7|Peruntukan|tempoh simpanan|"
     r"Syarat Kontrak|Design Guarantee|status|diserah|lewat|penyerahan|kelulusan|Skop subkon|pendaftaran ST|Detailed Work Program|"
     r"Sejarah revisi|Lampiran\b|akan disediakan|kebolehgunaan|Penyerahan bermasa|pengesahan|Pelulus",
     "S.O. / UTM / dokumen kontrak / rekod KESB", "Pengurus Projek"),
    ("Lukisan & teknikal (RFI)", r"refrigeran|R410a|R32|saiz paip|butiran|P12|kod |kod C|F5|lintol|skematik|struktur|jurutera|"
     r"pelan angkat|lukisan|asbestos|templat|C26|kuantiti|BQ berharga|jenis panel|kabel kebakaran|ketinggian pelekap|kawasan F9|"
     r"cat berplumbum|sisa terjadual|bahan berbahaya|busbar|RFI|S-JK|keperluan dengan|raised floor|paras lantai",
     "S.O. / perunding melalui RFI (KESB-QF-16)", "Pengurus Projek"),
    ("Kriteria ujian & standard", r"tekanan|vakum|suhu|ketebalan|kelembapan|flood|takungan|nilai|standard|kekerapan|had |"
     r"kerataan|penyaksian|slump|kiub|torque|ujian|makmal|data pengilang|kriteria",
     "S.O. / data pengilang / JKR 20800-0257-25", "Wakil Kualiti"),
    ("Lokasi & pengguna", r"wakil pengguna|fasa|ruang|notis|pengguna|kebenaran UTM|makluman",
     "UTM / wakil fakulti", "Pengurus Projek"),
    ("Loji & peralatan", r"loji|no\. siri|bilangan unit|janakuasa|alat kuasa|coring|penguji|pejabat tapak|kalibrasi|12 bulan",
     "Site Supervisor / SSS", "Site Supervisor"),
    ("Personel & organisasi", r"nama|lantikan|Wakil|PW4|ketua|CSS|juruaudit|penyemak|pemilik skop|Pengarah|kehadiran|"
     r"Pegawai Penguasa|Diluluskan oleh|pembekal|pemasang|Organisasi|Director|peranan",
     "Keputusan dalaman KESB", "Pengurus Projek"),
    ("Sisa & alam sekitar", r"sisa|RoRo|besi buruk|buangan|pelupusan",
     "UTM (Pengurusan Fasiliti) / S.O.", "Pengurus Projek"),
]


def kategori(text):
    for k, rx, src, own in KATEGORI:
        if re.search(rx, text, re.I):
            return k, src, own
    return "Lain-lain", "[tentukan]", "Pengurus Projek"


items = {}  # kunci -> dict


def add(desc, where, context=""):
    desc, context = desc.strip().rstrip("."), context.strip()
    # perkara ringkas / nombor sahaja → gabung dengan konteks supaya bermakna
    if context and (not desc or len(desc) < 45) and context.lower() not in desc.lower():
        desc = f"{context}: {desc}" if desc else context
    desc = desc or "Maklumat belum disahkan"
    key = re.sub(r"\W+", " ", desc.lower()).strip()
    it = items.setdefault(key, {"perkara": desc, "lokasi": [], "konteks": context.strip()})
    if where not in it["lokasi"]:
        it["lokasi"].append(where)


def scan_md(path, label):
    sec = ""
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if m := re.match(r"^(#{1,3})\s+(.*)", line):
            sec = m.group(2)[:40]
        for m in TAG.finditer(line):
            if line.startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                # sel pertama yang bukan nombor sahaja (cth jadual bernombor)
                ctx = next((c for c in cells[:3] if c and not re.fullmatch(r"[\d./]+", c)), "")
                ctx = re.sub(r"\*\*|\[PERLU SAHKAN[^\]]*\]", "", ctx).strip()[:70]
                if not m.group(1):
                    # tanda kosong: guna teks sebelum tanda dalam sel yang sama
                    cell = next((c for c in cells if m.group(0) in c), "")
                    before = re.sub(r"\*\*", "", cell[:cell.find(m.group(0))]).strip(" ;:–-")
                    if before and before not in ctx:
                        ctx = f"{ctx} – {before[-60:]}" if ctx else before[-60:]
                if sec.startswith("Sejarah revisi"):
                    ctx = "Pelulus CP-QC Rev 0 (Sejarah revisi)"
            else:
                before = TAG.sub("", line[:m.start()])
                before = re.sub(r"^[-\d.)\s*]+", "", before).replace("**", "").strip()
                ctx = before[-70:] if before else sec
            add(m.group(1), f"{label} – {sec}" if sec else label, ctx)


scan_md("output/CP-QC-draf.md", "CP-QC")
for f in sorted(Path("output/method-statement").glob("MS-*.md")):
    scan_md(f, f.name[:5])

d = json.loads(Path("output/itp.json").read_text(encoding="utf-8"))
for ti, t in enumerate(d["itp"], start=1):
    for ai, a in enumerate(t["aktiviti"], start=1):
        for v in a.values():
            for m in TAG.finditer(str(v)):
                add(m.group(1), f"ITP {ti}.{ai}", a["aktiviti"][:70])
for u in d["ujian_bahan"]:
    for k, v in u.items():
        for m in TAG.finditer(str(v)):
            add(m.group(1) or f"{u['bahan']} – {k}", f"ITP Ujian Bahan – {u['bahan'][:30]}", u["bahan"])

for m in TAG.finditer(Path("scripts/build_borang.py").read_text(encoding="utf-8")):
    add("Had selamat ujian udara ruang terkurung" if not m.group(1) else m.group(1), "Borang KESB-QF-18")

for desc in ["Bilangan unit & no. siri setiap loji / peralatan (Lampiran E)",
             "Kekerapan pemeriksaan berkala loji (Lampiran E)",
             "Kekerapan kalibrasi peralatan ukur & ujian (Lampiran E; cadangan 12 bulan)"]:
    add(desc, "Lampiran E – KESB/UTM-P3/PLT/001")

# --- Excel
HEAD = PatternFill("solid", fgColor="1F4E78")
CAT = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

rows = []
for it in items.values():
    k, src, own = kategori(it["perkara"] + " " + it["konteks"])
    rows.append((k, it, src, own))
order = [k for k, *_ in KATEGORI] + ["Lain-lain"]
rows.sort(key=lambda r: (order.index(r[0]), r[1]["perkara"].lower()))

wb = Workbook()
ws = wb.active
ws.title = "Item Terbuka"
ws.append(["Daftar Item Terbuka – maklumat [PERLU SAHKAN] (Lampiran J, CP-QC KESB/UTM-P3/CP-QC/001)"])
ws["A1"].font = Font(bold=True, size=13)
ws.append(["UTM – Baikpulih Makmal P&P, Pakej 3 – No. kontrak TN/2026-21/D-1/SP/BPA. Dijana daripada dokumen projek; "
           "kemas kini dokumen sumber kemudian jana semula (py -3.12 scripts/build_item_terbuka.py)."])
ws.append([])
head = ["No", "Kategori", "Perkara untuk disahkan", "Lokasi dalam dokumen", "Dapatkan daripada", "Pemilik",
        "Tarikh sasaran", "Status", "Jawapan / catatan"]
ws.append(head)
for c in ws[4]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEAD
    c.alignment = WRAP
    c.border = BORDER
prev = None
n = 0
for k, it, src, own in rows:
    if k != prev:
        ws.append(["", k.upper()])
        for c in ws[ws.max_row]:
            c.fill = CAT
            c.border = BORDER
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        prev = k
    n += 1
    ws.append([f"T-{n:03d}", k, it["perkara"], "; ".join(it["lokasi"]), src, own, "", "Terbuka", ""])
    for c in ws[ws.max_row]:
        c.alignment = WRAP
        c.border = BORDER
for col, w in zip("ABCDEFGHI", [8, 20, 55, 40, 28, 18, 13, 10, 35]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C5"
ws.auto_filter.ref = f"A4:I{ws.max_row}"

# --- Ringkasan untuk HQ (sheet pertama)
from collections import Counter

cnt = Counter(r[0] for r in rows)
hq = wb.create_sheet("Ringkasan untuk HQ", 0)
hq.column_dimensions["A"].width = 6
hq.column_dimensions["B"].width = 60
hq.column_dimensions["C"].width = 22
hq.column_dimensions["D"].width = 45


def put(values, bold=False, fill=None):
    hq.append(values)
    for c in hq[hq.max_row]:
        c.alignment = WRAP
        if bold:
            c.font = Font(bold=True, color="FFFFFF" if fill is HEAD else None)
        if fill:
            c.fill = fill
        if values and any(values):
            c.border = BORDER


hq.append(["Ringkasan untuk HQ KESB – Draf Construction Plan (Quality Control) Rev 0"])
hq["A1"].font = Font(bold=True, size=14)
hq.append(["UTM – Baikpulih Makmal P&P, Pakej 3 – No. kontrak TN/2026-21/D-1/SP/BPA. Disediakan oleh: Muhamad Azrin bin Rahmat, Pengurus Projek."])
hq.append([])
put(["", "A. KANDUNGAN PAKEJ DRAF", "Status", "Fail"], True, HEAD)
for item, st, f in [
    # Nama fail PDF ikut pakej serahan HQ (scripts/build_pakej_hq.py); fail Word & Excel asal dalam subfolder "Fail Asal (Word & Excel)"
    ("Memo iringan kepada HQ KESB", "–", "00 Memo Iringan kepada HQ.pdf"),
    ("Construction Plan (Quality Control) – KESB/UTM-P3/CP-QC/001", "Draf Rev 0", "01 Construction Plan (QC) Rev0.pdf"),
    ("Lampiran A: Inspection & Test Plan (ITP) + daftar NCR, dokumen, lokasi, log isu", "Draf Rev 0", "02 Lampiran A - ITP.pdf"),
    ("Lampiran B: Carta organisasi projek", "Draf 22/10/2026", "03 Lampiran B - Carta Organisasi.pdf"),
    ("Lampiran C: Daftar lukisan (264 helaian)", "Lukisan tender", "04 Lampiran C - Daftar Lukisan.pdf"),
    ("Lampiran E: Jadual loji – KESB/UTM-P3/PLT/001 (Prelim 3.0 (iv))", "Draf – butiran loji belum diisi", "05 Lampiran E - Jadual Loji.pdf"),
    ("Lampiran H: Method statement MS-01 – MS-14 (Prelim 3.0 (ii))", "Draf Rev 0", "06 Lampiran H - Method Statement MS-01 hingga MS-14.pdf"),
    ("Lampiran I: Borang KESB-QF-01 – QF-29", "Rev 0", "07 Lampiran I - Borang QF-01 hingga QF-29.pdf"),
    ("Lampiran J: Daftar Item Terbuka (fail ini)", "Dijana daripada dokumen", "08 Lampiran J - Daftar Item Terbuka.pdf"),
    ("Lampiran K: Carta alir sistem kualiti (K0–K9)", "Rev 0", "09 Lampiran K - Carta Alir.pdf"),
]:
    put(["", item, st, f])
hq.append([])
put(["", "B. ITEM TERBUKA MENGIKUT KATEGORI", "Bil. item", "Dapatkan daripada"], True, HEAD)
for k, _, src, _ in KATEGORI:
    if cnt.get(k):
        put(["", k, cnt[k], src])
if cnt.get("Lain-lain"):
    put(["", "Lain-lain", cnt["Lain-lain"], "[tentukan]"])
put(["", "JUMLAH", n, ""], True, CAT)
hq.append([])
put(["", "C. ITEM KRITIKAL UNTUK PERHATIAN HQ", "Kesan", "Tindakan dicadang"], True, HEAD)
for item, kesan, tindakan in [
    ("Dokumen wajib Prelim 3.0 (i)–(iv) belum/lewat diserah (tarikh akhir 24/08 atau 28/08/2026)", "Pelanggaran syarat kontrak",
     "Lengkapkan Lampiran E & programme; serah segera kepada S.O."),
    ("Pendaftaran Suruhanjaya Tenaga KESB tamat 28/09/2026; sijil MOF tamat 14/03/2026", "Kelayakan kerja elektrik / perolehan",
     "Sahkan status pembaharuan; sahkan pendaftaran ST H&Z"),
    ("Kelulusan UTM bagi subkon (Harta Perkasa, H&Z) & Deed of Assignment (SST para 9(g), 12)", "UTM berhak batal SST",
     "Sahkan status kelulusan; mohon jika belum"),
    ("Lampiran A SST (LAD, DLP) belum diterima", "Liabiliti LAD & tempoh DLP tidak diketahui", "Dapatkan salinan lengkap SST (hlm 7–15)"),
    ("Detailed Work Program siap 16/08/2028 – melepasi tarikh siap kontrak 14/08/2028", "Programme tidak sah", "Kemas kini programme"),
    ("Wakil Kualiti, Site Agent tambahan, Orang Kompeten PW4, penyelia elektrik & mekanikal belum dinamakan", "Klausa 59; BQ E nota IV",
     "Keputusan lantikan HQ"),
    ("Isu lukisan untuk RFI: S-FK21 & E-FKM32 tiada, tiada skematik E04–P23, R32/R410a, P12", "Kerja tidak boleh dimulakan di lokasi terlibat",
     "Keluarkan RFI kepada S.O."),
    ("Wakil pengguna (nama & no. telefon) dan fasa setiap ruang", "Koordinasi lokasi tidak boleh dilaksana", "Minta daripada UTM melalui S.O."),
]:
    put(["", item, kesan, tindakan])
hq.append([])
hq.append(["", "Senarai penuh setiap item, lokasinya dalam dokumen dan pemilik: sheet \"Item Terbuka\"."])

out = Path("output/Daftar-Item-Terbuka.xlsx")
wb.save(out)
print(f"Siap: {out}  ({n} item)")
for k, c in Counter(r[0] for r in rows).items():
    print(f"  {k}: {c}")
