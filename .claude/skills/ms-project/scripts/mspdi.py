"""Baca & analisis fail MS Project XML (MSPDI). Hanya stdlib + openpyxl.

Fail .mpp perlu ditukar dahulu dengan mpp_convert.ps1 (-Out x.xml).

Subarahan:
  summary   XML [--contract-finish YYYY-MM-DD]       ringkasan projek & tarikh siap
  tasks     XML --out tasks.xlsx|.csv                senarai penuh tugasan
  critical  XML [--out x.xlsx]                       tugasan kritikal (slack <= 0)
  lookahead XML --from YYYY-MM-DD [--weeks 3] [--out x.xlsx]
  late      XML [--status-date YYYY-MM-DD] [--out x.xlsx]   lewat berbanding baseline
  scurve    XML --out x.xlsx [--period month|week] [--weight cost|duration] [--status-date ...]
  check     XML [--out x.xlsx]                       semakan kualiti jadual

Semua tarikh dalam output ialah YYYY-MM-DD. Tempoh dalam hari bekerja
(ikut MinutesPerDay projek). Kos dalam RM.
"""
import argparse, csv, sys, xml.etree.ElementTree as ET
from datetime import datetime, date, timedelta

NS = {"p": "http://schemas.microsoft.com/project"}
LINK = {"0": "FF", "1": "FS", "2": "SF", "3": "SS"}
CONSTRAINT = {"0": "ASAP", "1": "ALAP", "2": "MSO", "3": "MFO", "4": "SNET", "5": "SNLT", "6": "FNET", "7": "FNLT"}


def _t(el, tag, default=None):
    e = el.find("p:" + tag, NS)
    return e.text if e is not None and e.text is not None else default


def _dt(s):
    return datetime.fromisoformat(s) if s else None


def _hours(dur):
    """PT5840H30M0S -> jam (float)."""
    if not dur or not dur.startswith("PT"):
        return 0.0
    s = dur[2:]
    h = m = sec = 0.0
    for unit in ("H", "M", "S"):
        if unit in s:
            v, s = s.split(unit, 1)
            v = float(v or 0)
            if unit == "H": h = v
            elif unit == "M": m = v
            else: sec = v
    return h + m / 60 + sec / 3600


def load(path):
    root = ET.parse(path).getroot()
    mpd = int(_t(root, "MinutesPerDay", "480")) / 60.0
    proj = {
        "title": _t(root, "Title") or _t(root, "Name"),
        "start": _dt(_t(root, "StartDate")),
        "finish": _dt(_t(root, "FinishDate")),
        "status_date": _dt(_t(root, "StatusDate")),
        "hours_per_day": mpd,
        "currency": _t(root, "CurrencySymbol", "RM"),
    }
    tasks = []
    for el in root.findall("p:Tasks/p:Task", NS):
        if _t(el, "IsNull") == "1":
            continue
        bl = next((b for b in el.findall("p:Baseline", NS) if _t(b, "Number") == "0"), None)
        preds = []
        for pl in el.findall("p:PredecessorLink", NS):
            lag_days = int(_t(pl, "LinkLag", "0")) / 10 / 60 / mpd
            preds.append((_t(pl, "PredecessorUID"), LINK.get(_t(pl, "Type", "1"), "FS"), lag_days))
        t = {
            "uid": _t(el, "UID"), "id": int(_t(el, "ID", "0")), "wbs": _t(el, "WBS", ""),
            "outline": _t(el, "OutlineNumber", ""), "level": int(_t(el, "OutlineLevel", "0")),
            "name": (_t(el, "Name", "") or "").strip(),
            "summary": _t(el, "Summary") == "1", "milestone": _t(el, "Milestone") == "1",
            "critical": _t(el, "Critical") == "1", "manual": _t(el, "Manual") == "1",
            "active": _t(el, "Active", "1") == "1",
            "start": _dt(_t(el, "Start")), "finish": _dt(_t(el, "Finish")),
            "dur": _hours(_t(el, "Duration")) / mpd,
            "pct": int(_t(el, "PercentComplete", "0")),
            "act_start": _dt(_t(el, "ActualStart")), "act_finish": _dt(_t(el, "ActualFinish")),
            "total_slack": int(_t(el, "TotalSlack", "0")) / 10 / 60 / mpd,
            "cost": int(float(_t(el, "Cost", "0"))) / 100.0,
            "constraint": CONSTRAINT.get(_t(el, "ConstraintType", "0"), "?"),
            "constraint_date": _dt(_t(el, "ConstraintDate")),
            "bl_start": _dt(_t(bl, "Start")) if bl is not None else None,
            "bl_finish": _dt(_t(bl, "Finish")) if bl is not None else None,
            "bl_cost": int(float(_t(bl, "Cost", "0"))) / 100.0 if bl is not None else 0.0,
            "preds": preds,
        }
        tasks.append(t)
    by_uid = {t["uid"]: t for t in tasks}
    for t in tasks:
        t["succ"] = []
    for t in tasks:
        for puid, typ, lag in t["preds"]:
            if puid in by_uid:
                by_uid[puid]["succ"].append(t["uid"])
    # nama induk (ringkasan) untuk konteks
    stack = {}
    for t in sorted(tasks, key=lambda x: x["id"]):
        stack[t["level"]] = t["name"]
        t["parent"] = stack.get(t["level"] - 1, "")
        t["path"] = " > ".join(stack[l] for l in range(1, t["level"]) if l in stack)
    return proj, tasks, by_uid


def d(x):
    return x.strftime("%Y-%m-%d") if x else ""


def leaves(tasks):
    return [t for t in tasks if not t["summary"] and t["active"] and t["id"] > 0]


def pred_text(t, by_uid):
    out = []
    for puid, typ, lag in t["preds"]:
        p = by_uid.get(puid)
        s = str(p["id"]) if p else "UID" + str(puid)
        if typ != "FS" or lag:
            s += typ
        if lag:
            s += ("+" if lag > 0 else "") + f"{lag:g}d"
        out.append(s)
    return ", ".join(out)


# ---------------- output ----------------
COLS = [("ID", "id"), ("WBS", "wbs"), ("Nama", "name"), ("Induk", "parent"), ("Ringkasan", "summary"),
        ("Tempoh (hari)", "dur"), ("Mula", "start"), ("Siap", "finish"), ("Mula BL", "bl_start"),
        ("Siap BL", "bl_finish"), ("% Siap", "pct"), ("Mula Aktual", "act_start"), ("Siap Aktual", "act_finish"),
        ("Total Slack (hari)", "total_slack"), ("Kritikal", "critical"), ("Predecessor", "_pred"),
        ("Kekangan", "constraint"), ("Tarikh Kekangan", "constraint_date"), ("Kos (RM)", "cost")]


def row(t, by_uid, cols):
    out = []
    for _, k in cols:
        v = pred_text(t, by_uid) if k == "_pred" else t.get(k)
        if isinstance(v, datetime): v = d(v)
        elif isinstance(v, bool): v = "Ya" if v else ""
        elif isinstance(v, float): v = round(v, 2)
        out.append(v)
    return out


def write(rows, headers, out, title=None, extra=None):
    """Tulis ke .xlsx / .csv, atau cetak ke skrin bila out kosong."""
    if not out:
        print(" | ".join(headers))
        for r in rows:
            print(" | ".join("" if v is None else str(v) for v in r))
        return
    if out.lower().endswith(".csv"):
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f); w.writerow(headers); w.writerows(rows)
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = (title or "Data")[:31]
        r0 = 1
        if extra:
            for line in extra:
                ws.cell(row=r0, column=1, value=line).font = Font(bold=(r0 == 1)); r0 += 1
            r0 += 1
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=r0, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            r0 += 1
            for c, v in enumerate(r, 1):
                ws.cell(row=r0, column=c, value=v)
        for c, h in enumerate(headers, 1):
            w = 50 if h in ("Nama", "Induk", "Isu") else 14
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
        ws.freeze_panes = ws.cell(row=r0 - len(rows) + 0, column=1) if rows else None
        ws.auto_filter.ref = f"A{r0 - len(rows)}:{ws.cell(row=r0, column=len(headers)).coordinate}" if rows else None
        wb.save(out)
    print(f"OK: {out} ({len(rows)} baris)")


# ---------------- subarahan ----------------
def cmd_summary(a):
    proj, tasks, by_uid = load(a.xml)
    lv = leaves(tasks)
    top = [t for t in tasks if t["id"] == 0] or [t for t in tasks if t["level"] == 1]
    total_cost = sum(t["cost"] for t in lv)
    earned = sum(t["cost"] * t["pct"] / 100 for t in lv)
    days = (proj["finish"].date() - proj["start"].date()).days + 1
    print(f"Projek       : {proj['title']}")
    print(f"Mula / Siap  : {d(proj['start'])} -> {d(proj['finish'])} ({days} hari kalendar)")
    print(f"Status date  : {d(proj['status_date']) or '(tiada)'}")
    if a.contract_finish:
        cf = date.fromisoformat(a.contract_finish)
        diff = (proj["finish"].date() - cf).days
        verdict = "LEWAT" if diff > 0 else "dalam tempoh"
        print(f"Siap kontrak : {cf} -> jadual {verdict} ({diff:+d} hari)")
    bl_fin = max((t["bl_finish"] for t in lv if t["bl_finish"]), default=None)
    print(f"Siap baseline: {d(bl_fin) or '(tiada baseline)'}")
    print(f"Tugasan      : {len(tasks)} (aktiviti {len(lv)}, ringkasan {sum(t['summary'] for t in tasks)}, "
          f"milestone {sum(t['milestone'] for t in lv)})")
    print(f"Kritikal     : {sum(t['critical'] for t in lv)} aktiviti")
    print(f"Kos dimuat   : {proj['currency']}{total_cost:,.2f}")
    if total_cost:
        print(f"Kemajuan (kos): {earned / total_cost * 100:.2f}%  ({proj['currency']}{earned:,.2f})")
    print("\nPecahan aras 2:")
    for t in [t for t in tasks if t["level"] == 2]:
        print(f"  {t['wbs']:<6} {t['name'][:60]:<60} {d(t['start'])} -> {d(t['finish'])}  {t['pct']:>3}%  RM{t['cost']:,.0f}")


def cmd_tasks(a):
    proj, tasks, by_uid = load(a.xml)
    write([row(t, by_uid, COLS) for t in tasks if t["id"] > 0], [h for h, _ in COLS], a.out, "Tugasan")


def cmd_critical(a):
    proj, tasks, by_uid = load(a.xml)
    sel = [t for t in leaves(tasks) if t["critical"] or t["total_slack"] <= 0]
    sel.sort(key=lambda t: (t["start"], t["id"]))
    cols = [c for c in COLS if c[1] in ("id", "wbs", "name", "parent", "dur", "start", "finish", "total_slack", "pct", "_pred")]
    write([row(t, by_uid, cols) for t in sel], [h for h, _ in cols], a.out, "Kritikal")


def cmd_lookahead(a):
    proj, tasks, by_uid = load(a.xml)
    f = datetime.fromisoformat(a.from_)
    to = f + timedelta(weeks=a.weeks)
    sel = [t for t in leaves(tasks) if t["pct"] < 100 and t["start"] < to and t["finish"] >= f]
    sel.sort(key=lambda t: (t["start"], t["id"]))
    cols = [c for c in COLS if c[1] in ("id", "wbs", "name", "parent", "dur", "start", "finish", "pct", "critical", "_pred")]
    write([row(t, by_uid, cols) for t in sel], [h for h, _ in cols], a.out, "Lookahead",
          [f"LOOKAHEAD {a.weeks} MINGGU: {d(f)} hingga {d(to - timedelta(days=1))}", proj["title"]])


def _planned_pct(t, at):
    s, e = t["bl_start"] or t["start"], t["bl_finish"] or t["finish"]
    if at <= s: return 0.0
    if at >= e: return 100.0
    return (at - s).total_seconds() / max((e - s).total_seconds(), 1) * 100


def cmd_late(a):
    proj, tasks, by_uid = load(a.xml)
    sd = datetime.fromisoformat(a.status_date) if a.status_date else proj["status_date"]
    if not sd:
        sys.exit("Tiada status date dalam fail. Guna --status-date YYYY-MM-DD.")
    sd = sd.replace(hour=23, minute=59)
    rows = []
    for t in leaves(tasks):
        if t["pct"] >= 100: continue
        plan = _planned_pct(t, sd)
        if plan <= t["pct"]: continue
        isu = []
        if (t["bl_start"] or t["start"]) <= sd and t["pct"] == 0: isu.append("Belum mula")
        if (t["bl_finish"] or t["finish"]) <= sd: isu.append("Sepatutnya siap")
        slip = (t["finish"] - t["bl_finish"]).days if t["bl_finish"] else 0
        rows.append([t["id"], t["wbs"], t["name"], t["parent"], d(t["bl_start"]), d(t["bl_finish"]),
                     round(plan), t["pct"], round(plan - t["pct"]), slip, "Ya" if t["critical"] else "",
                     "; ".join(isu) or "Kemajuan rendah"])
    rows.sort(key=lambda r: (-r[8]))
    write(rows, ["ID", "WBS", "Nama", "Induk", "Mula BL", "Siap BL", "% Rancang", "% Aktual", "Beza %",
                 "Anjakan siap (hari)", "Kritikal", "Isu"], a.out, "Lewat",
          [f"AKTIVITI LEWAT BERBANDING BASELINE – status {d(sd)}", proj["title"]])


def _periods(start, end, period):
    p = date(start.year, start.month, 1) if period == "month" else start.date() - timedelta(days=start.weekday())
    out = []
    while p <= end.date():
        nxt = (date(p.year + (p.month == 12), p.month % 12 + 1, 1) if period == "month" else p + timedelta(days=7))
        out.append((p, nxt - timedelta(days=1)))
        p = nxt
    return out


def cmd_scurve(a):
    proj, tasks, by_uid = load(a.xml)
    lv = [t for t in leaves(tasks) if (t["bl_start"] or t["start"])]
    wkey = (lambda t: t["bl_cost"] or t["cost"]) if a.weight == "cost" else (lambda t: max(t["dur"], 0.0))
    total = sum(wkey(t) for t in lv)
    if not total:
        sys.exit("Tiada pemberat (kos/tempoh) – cuba --weight duration.")
    sd = datetime.fromisoformat(a.status_date) if a.status_date else (proj["status_date"] or datetime.now())
    s = min(t["bl_start"] or t["start"] for t in lv)
    e = max(max(t["bl_finish"] or t["finish"], t["finish"]) for t in lv)
    rows = []
    for p0, p1 in _periods(s, e, a.period):
        at = datetime.combine(p1, datetime.max.time())
        planned = sum(wkey(t) * _planned_pct(t, at) / 100 for t in lv) / total * 100
        # jadual semasa (forecast): guna Start/Finish semasa
        fc = sum(wkey(t) * (0 if at <= t["start"] else 1 if at >= t["finish"] else
                            (at - t["start"]).total_seconds() / max((t["finish"] - t["start"]).total_seconds(), 1))
                 for t in lv) / total * 100
        actual = round(sum(wkey(t) * t["pct"] / 100 for t in lv) / total * 100, 2) if p0 <= sd.date() else None
        rows.append([p1.strftime("%Y-%m-%d"), p0.strftime("%b %Y") if a.period == "month" else f"Mgu {p0}",
                     round(planned, 2), round(fc, 2), actual])
    # % aktual hanya diketahui pada status date – kosongkan selain tempoh semasa
    cur = [i for i, r in enumerate(rows) if r[4] is not None]
    for i in cur[:-1]:
        rows[i][4] = None
    headers = ["Akhir tempoh", "Tempoh", "% Rancang (baseline)", "% Forecast (jadual semasa)", "% Aktual (status)"]
    write(rows, headers, a.out, "S-Curve",
          [f"S-CURVE ({'kos' if a.weight == 'cost' else 'tempoh'}) – status {d(sd)}", proj["title"]])
    if a.out and a.out.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        from openpyxl.chart import LineChart, Reference
        wb = load_workbook(a.out); ws = wb.active
        hdr = 4
        n = hdr + len(rows)
        ch = LineChart(); ch.title = "S-Curve"; ch.y_axis.title = "% kumulatif"; ch.x_axis.title = "Tempoh"
        ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.height = 10; ch.width = 24
        ch.add_data(Reference(ws, min_col=3, max_col=5, min_row=hdr, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=hdr + 1, max_row=n))
        ws.add_chart(ch, "G4"); wb.save(a.out)


def cmd_check(a):
    proj, tasks, by_uid = load(a.xml)
    rows = []
    def add(t, isu, tahap):
        rows.append([t["id"], t["wbs"], t["name"], t["parent"], tahap, isu])
    lv = leaves(tasks)
    has_bl = any(t["bl_start"] for t in lv)
    if not has_bl:
        rows.append(["", "", "(projek)", "", "Tinggi", "Tiada baseline disimpan – S-curve rancangan & analisis lewat tidak tepat"])
    for t in lv:
        if t["milestone"]:
            pass
        elif not t["preds"] and t["start"] > proj["start"]:
            add(t, "Tiada predecessor (tergantung di awal)", "Sederhana")
        if not t["succ"] and t["finish"] < proj["finish"] - timedelta(days=1):
            add(t, "Tiada successor (tergantung di akhir)", "Sederhana")
        if t["constraint"] not in ("ASAP", "ALAP"):
            add(t, f"Kekangan {t['constraint']} {d(t['constraint_date'])}", "Rendah")
        if t["total_slack"] < 0:
            add(t, f"Slack negatif {t['total_slack']:.0f} hari", "Tinggi")
        if t["manual"]:
            add(t, "Tugasan dijadual manual", "Sederhana")
        if t["dur"] > 44 and not t["milestone"] and t["level"] > 2:
            add(t, f"Tempoh panjang {t['dur']:.0f} hari (> 2 bulan) – pertimbang pecahkan", "Rendah")
        if any(lag < 0 for _, _, lag in t["preds"]):
            add(t, "Lag negatif (lead)", "Rendah")
        if t["pct"] > 0 and not t["act_start"]:
            add(t, "% siap tanpa Mula Aktual", "Sederhana")
    order = {"Tinggi": 0, "Sederhana": 1, "Rendah": 2}
    rows.sort(key=lambda r: (order.get(r[4], 3), r[0] or 0))
    write(rows, ["ID", "WBS", "Nama", "Induk", "Tahap", "Isu"], a.out, "Semakan",
          [f"SEMAKAN KUALITI JADUAL – {proj['title']}"])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    def P(name, fn):
        p = sp.add_parser(name); p.add_argument("xml"); p.set_defaults(fn=fn); return p
    P("summary", cmd_summary).add_argument("--contract-finish")
    P("tasks", cmd_tasks).add_argument("--out", required=True)
    P("critical", cmd_critical).add_argument("--out")
    p = P("lookahead", cmd_lookahead); p.add_argument("--from", dest="from_", required=True)
    p.add_argument("--weeks", type=int, default=3); p.add_argument("--out")
    p = P("late", cmd_late); p.add_argument("--status-date"); p.add_argument("--out")
    p = P("scurve", cmd_scurve); p.add_argument("--out", required=True)
    p.add_argument("--period", choices=["month", "week"], default="month")
    p.add_argument("--weight", choices=["cost", "duration"], default="cost"); p.add_argument("--status-date")
    P("check", cmd_check).add_argument("--out")
    sys.stdout.reconfigure(encoding="utf-8")
    a = ap.parse_args(); a.fn(a)


if __name__ == "__main__":
    main()
